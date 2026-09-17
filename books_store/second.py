#----------------    BACKEND    -----------------

#------------------------------------------------
# IMPORTANTS IMPORTS FOR THIS PROGRAM
#------------------------------------------------

import openpyxl
from colorama import init, Fore, Back, Style
init(autoreset=True)
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font,Border,Side
from rich import console
from datetime import datetime, timedelta
from pyfiglet import figlet_format
from pyfiglet import Figlet
import pyfiglet

import csv
import random
import time
import uuid
import os
import emoji

from getpass import getpass
from datetime import datetime, timedelta
from pyfiglet import figlet_format

from rich.console import Console,Group
from rich.table import Table
from rich.text import Text
from rich.align import Align
from rich.rule import Rule
from rich import box
from rich.panel import Panel
from rich.progress import track
console = Console()
from function_utils import *

#------------------------------------------------
# ASSIGNING IMP PATHS TO THE VARIABLES
#------------------------------------------------

books_data_path = r"C:\Users\HP\Desktop\training\Python\Books_Store_Project\books_store\books_copy.xlsx"
users_data_path = r"C:\Users\HP\Desktop\training\Python\LIB\users_data.xlsx"
empls_data_path = r"C:\Users\HP\Desktop\training\Python\LIB\empls_data.csv"
credt_data_path = r"C:\Users\HP\Desktop\training\Python\LIB\credentials.txt"

#------------------------------------------------
# PROGRAM START FROM HERE :
#------------------------------------------------

def add_new_book(result = ""):


    if result == "new":
        head_color("\n--- Add New Genre ---\n")

        genre = user_input(f"Enter genre of {result} book : ").upper()
        decor_line()

        text_color("Add atleast one book's details under this new genre\n")

        new_code = user_input("Enter general unique code for books under this genre (ex. MYTH, NOV) : ").upper()
        decor_line()
                
        bookname = user_input(f"Enter proper name of {result} book : ")
        decor_line()

        authorname = user_input(f"Enter author name of {result} book : ")
        decor_line()
    
        lang = user_input(f"Enter language in which {result} book is written : ").upper()
        decor_line()

        publish_date = user_input(f"Enter publish date of {result} book in format 'dd-mm-yy': ")
        decor_line()

        while True:
            buy_price = user_input(f"Enter price of {result} book : ")
            decor_line()

            if buy_price.isdigit() and int(buy_price) < 0:
                error_message("Please enter valid price !")
            else:
                break

        while True:
            quantity = user_input(f"Enter quantity of {result} book : ")
            decor_line()

            if quantity.isdigit() and int(quantity) < 0:
                error_message("Please enter valid quantity !")
            else:
                break

        wb = load_workbook(books_data_path)
        s = wb.create_sheet(str(genre) + " -- " + str(new_code))
        header = ["Sr No","DDC Code","Book Name","Author Name","Language","Published Date","Wholesale Price","Market Price (INR)","Profit Margin","Quantities Available"]
        col_width = [7,15,50,28,17,20,18,21,20,25]
        f = Font(bold=True, underline="single")
        b = Border(left=Side("thin"),right=Side("thin"),top=Side("thin"),bottom=Side("thin"))
        pm = int(int(buy_price) * 30/100)
        wp = int(int(buy_price) - pm)

        for i in range(len(header)):
            col = i + 1
            c = s.cell(row=1, column=col)
            c.value = header[i]
            c.font = f
            c.alignment = Alignment(horizontal="center")
            c.border = b
            s.column_dimensions[c.column_letter].width = col_width[i]
        s.append([1,f"{new_code}10001",bookname,authorname,lang,publish_date,wp,int(buy_price),pm,int(quantity)])
        for cell in s[s.max_row]:
            cell.alignment = Alignment(horizontal='center', vertical='center')

        wb.save(books_data_path)

    else:

        head_color("\n--- Add New Book ---\n")
        display_genres()
        decor_line()
        genre = user_input(f"Enter exact genre code of desire book from above table : ").upper()
        decor_line()
        
        wd = load_workbook(books_data_path)
        sheet = None
        for s in wd.sheetnames:
            if genre in s:
                sheet = s
                break
        if sheet is None:
            error_message("Genre Not Found!")
            return
             
        bookname = user_input(f"Enter proper name of book : ")
        decor_line()

        authorname = user_input(f"Enter author name of book : ")
        decor_line()
    
        lang = user_input(f"Enter language in which book is written : ").upper()
        decor_line()

        publish_date = user_input(f"Enter publish date of book in format 'dd-mm-yy': ")
        decor_line()

        while True:
            buy_price = user_input(f"Enter price of book : ")
            decor_line()

            if buy_price.isdigit() and int(buy_price) < 0:
                error_message("Please enter valid price !")
            else:
                break

        while True:
            quantity = user_input(f"Enter quantity of book : ")
            decor_line()

            if quantity.isdigit() and int(quantity) < 0:
                error_message("Please enter valid quantity !")
            else:
                break

        ws = wd[sheet]
        last_row = ws.max_row
        last_sr = ws.cell(row=last_row,column=1).value
        last_uc = ws.cell(row=last_row,column=2).value
        new_sr = int(last_sr) + 1
        text,digt = "",""
        pm = int(int(buy_price) * 30/100)
        wp = int(int(buy_price) - pm)
        for i in last_uc:
            if i.isdigit() == False:
                text = text + i
            else:
                digt = digt + i
        new_digt = int(digt) + 1
        new_uc = text + str(new_digt)
        ws.append([new_sr,new_uc,bookname,authorname,lang,publish_date,wp,int(buy_price),pm,int(quantity)])

        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        wd.save(books_data_path)
        correct_message("Book Added Successfully!")


def add_copies():
    head_color("\n--- Add Copies Of Books ---\n")
    while True:
        decor_line()
        unique_code = user_input("Enter unique code of book : ").upper()
        wb = load_workbook(books_data_path)
        a = 0
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(min_row=2,values_only=False):
                if str(row[1].value).strip() == unique_code.upper():
                    a = 1
                    break

            if a == 1:
                break

        if a == 1:
            correct_message(" Book Found with this unique code ")

            while True:
                decor_line()
                quantity = user_input("Enter quantity of copies of this book to be added : ")

                if quantity.isdigit() and int(quantity) > 0:
                    row[7].value = int(row[7].value) + int(quantity)
                    wb.save(books_data_path)
                    correct_message("Data updated successfully ! ")
                    return
                else:
                    error_message(" Please enter valid quantity ! (ex. 2 or 4)")
        else:
            error_message("Book with this unique code not found !")

def add_old_book():
    decor_line()
    head_color("Actions Available :\n")
    text_color("\t1. Want to search Unique Code of book(s)?")
    text_color("\t2. Directly Add copies of book(s)")

    while True:
        decor_line()    
        sel_option = user_input("Enter option number from options : ")

        if sel_option == "1":
            search_book()
            break
        if sel_option == "2":
            add_copies()
            break
    

decor_line()
head_color("Actions Available :\n")
text_color("\t1. Search Book(s)")
text_color("\t2. Check Stock Balance")
text_color("\t3. Add Copies of Available Book(s)")
text_color("\t4. Add New Books(s)")
text_color("\t5. Add New Genre")
text_color("\t6. Exit")

while True:
    decor_line()    
    select_option = user_input("Enter option number from options : ")

    if select_option == "1":
        search_book()
        break

    elif select_option == "2":
        check_stock()
        break

    elif select_option == "3":
        add_old_book()
        break

    elif select_option == "4":
        add_new_book()
        break

    elif select_option == "5":
        add_new_book("new")
        break

    elif select_option == "6":
        break

    else:
        error_message("Please enter valid option nuber from above give options")
