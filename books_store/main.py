#----------------    BACKEND    -----------------

#------------------------------------------------
# IMPORTANTS IMPORTS FOR THIS PROGRAM
#------------------------------------------------

import openpyxl
from colorama import init, Fore, Back, Style
init(autoreset=True)
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font,Border,Side
from rich import console
from datetime import datetime, timedelta
from pyfiglet import figlet_format
from pyfiglet import Figlet
import pyfiglet



import csv
import random
import time
import uuid
import os
import emoji

from getpass import getpass
from datetime import datetime, timedelta
from pyfiglet import figlet_format

from rich.console import Console,Group
from rich.table import Table
from rich.text import Text
from rich.align import Align
from rich.rule import Rule
from rich import box
from rich.panel import Panel
from rich.progress import track
console = Console()
from function_utils import *
#------------------------------------------------
# ASSIGNING IMP PATHS TO THE VARIABLES
#------------------------------------------------

books_data_path = r"C:\Users\HP\Desktop\training\Python\Books_Store_Project\books_store\books_copy.xlsx"
users_data_path = r"C:\Users\HP\Desktop\training\Python\Books_Store_Project\books_store\users_data.xlsx"
empls_data_path = r"C:\Users\HP\Desktop\training\Python\Books_Store_Project\books_store\empls_data.csv"
credt_data_path = r"C:\Users\HP\Desktop\training\Python\Books_Store_Project\books_store\credentials.txt"
STORE_DATA_PATH = "STORE_RECORDS.xlsx"

#------------------------------------------------
# ASSIGNING VALUES TO IMP VARIABLES :
#------------------------------------------------


#------------------------------------------------
# DEFINING IMP FUNCTIONS OF DECORATIVE STUFFS :
#------------------------------------------------


#------------------------------------------------
# PROGRAM START FROM HERE :
#------------------------------------------------

# DEFINING IMP FUNCTIONS USED IN THIS PROGRAM ---

def first_p():
    def display_login_title():
        pass



        # while True:
        #     username = input("Enter Your Username : ")


    def create_store_records_excel():

        if not os.path.exists(STORE_DATA_PATH):

            workbook = openpyxl.Workbook()

            sheet1 = workbook.active
            sheet1.title = "Users Data"

            header1 = ["Unique Code", "Book Name", "Author Name", "MRP (in Rs)", "Quantity", "Total", "Profit", "Date"]
            column_width1 = [20,52,34,19,18,19,22,17]

            for i in range(len(header1)):
                col = i + 1
                c = sheet1.cell(row=1, column=col)
                c.value = header1[i]
                c.font = Font(bold=True, underline="single")
                c.alignment = Alignment(horizontal="center")
                c.border = Border(left=Side("thin"),right=Side("thin"),top=Side("thin"),bottom=Side("thin"))
                sheet1.column_dimensions[c.column_letter].width = column_width1[i]

            sheet2 = workbook.create_sheet(title="Store Data")

            header2 = ["Unique Code", "Book Name", "Author Name", "Wholesale Price", "Quantity", "Total Expense", "Date"]
            column_width2 = [20,52,34,19,18,22,17]


            for i in range(len(header2)):
                col = i + 1
                c = sheet2.cell(row=1, column=col)
                c.value = header2[i]
                c.font = Font(bold=True, underline="single")
                c.alignment = Alignment(horizontal="center")
                c.border = Border(left=Side("thin"),right=Side("thin"),top=Side("thin"),bottom=Side("thin"))
                sheet2.column_dimensions[c.column_letter].width = column_width2[i]

            workbook.save(STORE_DATA_PATH)
            workbook.close()

            # console.print("[green]Excel database created successfully![/green]")


    def create_bill(code,bookname,author,price,quantity,total,date):

        good_quotes = ["'Books are uniquely portable magic'",
                    "'Today a reader,tomorrow a leader'",
                    "'Read what you love until you love to read'"
                    "'There is no friend as loyal as a book'",
                    "'A book dream is a dream that you hold in your hand'"]
        quote = random.choice(good_quotes)

        Total = 0
        for i in total:
            Total = Total + int(i)
        decor_line()
        title_box = Table(box=box.ROUNDED, border_style="#FFF3E0",style="on #2d1f0f",expand=True,show_header=False,padding=(1,1))
        title_box.add_column(justify="center")
        title_box.add_row(Text("📚   DIGITAL LIBRARY OF NAVI MUMBAI   📚",style="#0066FF",justify="center"))
        title_box.add_row(Text("( 🌏 💻  By MJJ-TechWorld  💻 🌏 )",style="bold #FFD700 on #1a1a1a",justify="center"))

        table = Table(box=box.DOUBLE_EDGE,
                    title_style="bold yellow",
                    header_style="bold white on blue",
                    border_style="bright_yellow",
                    show_lines=True,
                    expand=True)
        for col in ["Sr.","Unique Code","Name Of Books","Author Name", "Price","Quantity","Total"]:
            table.add_column(col,justify="center",style="cyan",overflow="fold")

        for i in range(len(code)):
            table.add_row(str(i+1),str(code[i]),str(bookname[i]),str(author[i]),str(price[i]),str(quantity[i]),str(total[i]))

        content = Group(
            title_box,
            line("bright_cyan"),
            line("bright_magenta"),
            Text(f"Date : {date}", style="cyan", justify="right"),
            table,
            Text(f"Total : Rs.{Total}/-", style="bold green", justify="right"),
            line("bright_magenta"),
            Text("--- Thank You Visit Again ---", style="bold green", justify="center"),
            line("bright_magenta"),
            Text(f"{quote}", style="bold yellow", justify="center"),
            line("bright_magenta")
            )
        console.print(Panel(content, box=box.DOUBLE, border_style="#00BFFF on #0F2D2E",padding=(1,1), width=console.width - 2),"\n")

    def buy_book():
        code_list,bookname_list,author_list,price_list,quantity_list,total_list = [],[],[],[],[],[]
        date = datetime.today().strftime('%d-%m-%Y')
        while True:
            value = False
            decor_line()
            head_color("---Buy Book(s)---\n")
            uc = user_input("Enter Unique Code Of Book : ").upper()
            wb = load_workbook(books_data_path)
            for sheet in wb.sheetnames:
                s = wb[sheet]
                for row in s.iter_rows(min_row=2,values_only=True):
                    if uc == str(row[1]).strip():
                        correct_message(f"{row[2]} : By {row[3]} : Rs.{row[7]} : Avail. {row[9]}\n")
                        avail_q = row[9]
                        bookname = row[2]
                        author = row[3]
                        price = row[7]
                        value = True

            if not value:
                error_message("Book Not Found!") 
                continue


            while True:
                quantity = user_input("Enter quantity of this book : ")

                if quantity == "":
                    error_message("Quantity cannot be none!")
                    break

                if not quantity.isdigit():
                    error_message("Please enter valid quantity")
                    continue

                if int(quantity) < 0:
                    error_message("Please enter valid quantity!")
                    continue

                if int(quantity) <= int(avail_q):
                    correct_message("Book Sold Successfully!")

                    code_list.append(uc)
                    bookname_list.append(bookname)
                    author_list.append(author)
                    price_list.append(price)
                    quantity_list.append(quantity)
                    total_list.append(int(price)*int(quantity))
                    
                    wb = load_workbook(books_data_path)
                    for sheet in wb.sheetnames:
                        s = wb[sheet]
                        for row in s.iter_rows(min_row=2,values_only=False):
                            if row[1].value == uc:
                                row[9].value = int(row[9].value) - int(quantity)
                    wb.save(books_data_path)
                    break

                else:
                    error_message("Insufficient quantity of book available!")
                    continue

            while True:
                head_color("Want to buy more book(s)?\n")
                text_color("\t1.Yes")
                text_color("\t2.No")

                sel_option = ask_option_number()

                if sel_option == "1":
                    break

                elif sel_option == "2":
                    update_record(0,code_list,bookname_list,author_list,price_list,quantity_list,total_list,date)
                    create_bill(code_list,bookname_list,author_list,price_list,quantity_list,total_list,date)
                    return

                else:
                    option_error()
                    continue
            


    def choose_option():
        while True:
            decor_line()
            head_color("Available Options :\n")

            info_message("\t1. Want to check unique code of book once ?")
            info_message("\t2. Buy book(s) directly\n")

            sel_option = user_input("Enter option number to proceed further : ")

            if sel_option == "1":
                search_book()
                choose_option()
                return

            elif sel_option == "2":
                buy_book()
                menu_display()
                return
            else:
                error_message("Please enter valid option number from given options!")


    def change_rate():
        wb = load_workbook(books_data_path)
        for sheet in wb.sheetnames:
            s = wb[sheet]
            for row in s.iter_rows(min_row=2,values_only=False):
                if row[6].value is None:
                    continue
                row[6].value = int((int(row[7].value)) - int((int(row[7].value) * 30)/100))
                row[8].value = int((int(row[7].value) * 30)/100)
        wb.save(books_data_path)  


    def banner():
        f = Figlet(font='small',width=200)
        name = figlet_format('Digital Books\n Store',font="mini", width=200).strip()
        content = Group(
            Align.center(Text(name, style="bold bright_white on #00ACC1")),
            Text(""),
            Align.center(f.renderText("Digital Books Store"))
        )
        title_box = Panel(content,
                        box=box.ROUNDED,
                        border_style="#FF8C00",
                        style="on #2d1f0f",
                        padding=(1,2),
                        expand=False)

        banner_console = Console(width=120)
        banner_console.print(title_box)

        console.print(Panel(
            f"{name}\n[bold #FFD700 on #3a3200]  (By MJJ-TechWorld)  [/]",
            box=box.ROUNDED,
            border_style="#FF8C00",
            style="bold #B2EBF2 on #00363A",
            padding=(1, 3),
            expand=False
        ))


        # print(emoji.emojize(":globe_showing_Asia-Australia:"))
        # print(emoji.emojize(":open_book:"))
        # print(emoji.emojize(":books:"))
        # print(emoji.emojize(":laptop:"))

    def menu_display():
        create_store_records_excel()
        login_title()
        decor_line()
        head_color("Actions Available :\n")
        text_color("\t1. Search book(s)")
        text_color("\t2. Buy book(s)")
        text_color("\t3. Exit")

        while True:
            decor_line()    
            select_option = user_input("Enter option number from options : ")

            if select_option == "1":
                search_book()
                break
            elif select_option == "2":
                choose_option()
                break

            elif select_option == "3":
                break

            else:
                option_error()

    menu_display()



if __name__ == "__main__":
    first_p()






